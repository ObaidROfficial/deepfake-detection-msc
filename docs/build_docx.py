#!/usr/bin/env python3
"""
Assemble the MSc dissertation into a single .docx.

Arial 12, double-spaced, justified body. Automatic table of contents, list of
figures and list of tables (Word field codes - press Ctrl+A then F9 in Word to
populate them). Roman page numbers in the front matter, arabic from Chapter 1.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Mm, RGBColor

DOCS = Path(sys.argv[1])
FIGDIR = Path(sys.argv[2])
OUT = Path(sys.argv[3])
# Screenshots of the running artefact, kept apart from the generated plots.
SHOTDIR = DOCS.parent / "screenshots"

BODY_FONT = "Arial"

# --- figure placement -------------------------------------------------------
# Literal numbering rather than SEQ fields: the numbers are referenced by name
# throughout the prose, so they must not shift. Page numbers in the List of
# Figures still come from a TOC field over the caption style.
FIGURES = {
    "6.1 Test plan": [
        ("fig_5_01_class_balance.png", "Class balance in both datasets"),
        ("fig_5_02_samples_dataset_a.png", "Sample images from Dataset A (StyleGAN fakes, FFHQ reals)"),
        ("fig_5_03_samples_dataset_b.png", "Sample images from Dataset B (OpenForensics)"),
        ("fig_5_04_domain_gap.png", "Low-level statistical differences between the two datasets"),
        ("fig_5_05_class_separability.png", "Class separability of simple image statistics, Dataset A"),
        ("fig_5_06_colour_channels.png", "Colour channel distributions by class"),
    ],
    "6.2 In-domain results": [
        ("fig_6_01_baseline_roc.png", "ROC curves for both classical baselines, in-domain"),
        ("fig_6_cm_logistic_regression_A.png", "Confusion matrix, logistic regression, in-domain"),
        ("fig_6_cm_hog_lbp_svm_A.png", "Confusion matrix, HOG + LBP with linear SVC, in-domain"),
        ("fig_6_scores_logistic_regression.png", "Score distribution by class, logistic regression"),
        ("fig_6_scores_hog_lbp_svm.png", "Score distribution by class, HOG + LBP"),
        ("fig_6_cnn_curves_mobilenet_v2.png", "Training and validation curves, MobileNetV2"),
        ("fig_6_cnn_curves_efficientnet_b0.png", "Training and validation curves, EfficientNet-B0"),
        ("fig_6_cnn_cm_mobilenet_v2.png", "Confusion matrix, MobileNetV2, in-domain"),
        ("fig_6_cnn_cm_efficientnet_b0.png", "Confusion matrix, EfficientNet-B0, in-domain"),
    ],
    "6.3 Cross-dataset results": [
        ("fig_6_cm_logistic_regression_B.png", "Confusion matrix, logistic regression, cross-dataset"),
        ("fig_6_cm_hog_lbp_svm_B.png", "Confusion matrix, HOG + LBP, cross-dataset"),
        ("fig_6_cnn_scores_mobilenet_v2.png", "Score distributions in-domain against cross-dataset, MobileNetV2"),
        ("fig_6_cnn_scores_efficientnet_b0.png", "Score distributions in-domain against cross-dataset, EfficientNet-B0"),
        ("fig_6_curves_mobilenet_v2.png", "ROC and precision-recall curves, MobileNetV2"),
        ("fig_6_curves_efficientnet_b0.png", "ROC and precision-recall curves, EfficientNet-B0"),
        ("fig_6_26_spectral_profiles.png",
         "Radially averaged power spectra by class. The classes separate in "
         "Dataset B above 0.2 cycles per pixel and barely separate in Dataset A"),
    ],
    "6.5 Robustness results": [
        ("fig_5_07_jpeg_preview.png", "Visual effect of JPEG compression at the tested quality levels"),
        ("fig_5_08_noise_preview.png", "Visual effect of Gaussian noise at the tested magnitudes"),
        ("fig_7_01_robustness_mobilenet_v2.png", "Accuracy against perturbation magnitude, all families"),
        ("fig_7_02_worst_case_mobilenet_v2.png", "Worst-case accuracy by perturbation family"),
    ],
}
_n = 0
for _sec in FIGURES:
    lst = []
    for fn, cap in FIGURES[_sec]:
        _n += 1
        lst.append((fn, f"Figure 6.{_n}: {cap}"))
    FIGURES[_sec] = lst
TOTAL_FIGURES = _n

# Chapter 5 figures are screenshots of the artefact rather than plots, so they
# live in screenshots/ instead of results/figures/ and are numbered separately.
# Captions are deliberately short. Word builds the List of Figures from the
# whole caption paragraph, so a sentence-long caption makes that list
# unreadable. The explanation belongs in the body text.
CH4_FIGURES = {
    "4.4 Experimental design": [
        ("fig_4_01_threshold_diagnostic.jpg",
         "The both-polarity threshold diagnostic in evaluate.py"),
    ],
}

CH5_FIGURES = {
    "5.4 CNN implementation": [
        ("fig_5_01_freeze_unfreeze.jpg",
         "Model construction and the two-stage freeze/unfreeze in model.py"),
    ],
    "5.5 Demonstration tools": [
        ("fig_5_02_app_real_verdict.jpg",
         "The detector classifying a real photograph, with the out-of-range "
         "note. Image from Flickr-Faces-HQ (NVIDIA, CC BY-NC-SA 4.0)"),
        ("fig_5_03_app_fake_verdict.jpg",
         "The same interface on a StyleGAN face, which raises no warning"),
        ("fig_5_04_reliability_dialog.jpg",
         "The reliability panel, reporting measured accuracy by condition"),
        ("fig_5_05_out_of_range_checks.jpg",
         "The out-of-range checks in predict.py that produce those warnings"),
        ("fig_5_06_installer_options.jpg",
         "The installer's options screen"),
    ],
    "5.6 Robustness implementation": [
        ("fig_5_07_robustness_docstring.jpg",
         "The prediction recorded in robustness.py before the suite was run"),
        ("fig_5_08_fgsm_label_choice.jpg",
         "fgsm_evaluate in robustness.py, showing the label-choice argument"),
    ],
}
_m = 0
for _sec in CH4_FIGURES:
    lst = []
    for fn, cap in CH4_FIGURES[_sec]:
        _m += 1
        lst.append((fn, f"Figure 4.{_m}: {cap}"))
    CH4_FIGURES[_sec] = lst
TOTAL_FIGURES += _m
_m = 0
for _sec in CH5_FIGURES:
    lst = []
    for fn, cap in CH5_FIGURES[_sec]:
        _m += 1
        lst.append((fn, f"Figure 5.{_m}: {cap}"))
    CH5_FIGURES[_sec] = lst
TOTAL_FIGURES += _m

FILES = [
    ("ch1_introduction.md", "Chapter 1: Introduction"),
    ("ch2_literature_review.md", "Chapter 2: Literature Review and Fact Finding"),
    ("ch3_methodology.md", "Chapter 3: Requirements and Methodology"),
    ("ch4_design.md", "Chapter 4: Design"),
    ("ch5_implementation.md", "Chapter 5: Implementation"),
    ("ch6_results.md", "Chapter 6: Testing, Results and Evaluation"),
    ("ch7_critical_evaluation.md", "Chapter 7: Critical Evaluation"),
    ("ch8_conclusion.md", "Chapter 8: Conclusion and Further Work"),
]


# --- low-level helpers ------------------------------------------------------
def field(paragraph, instr: str, placeholder: str = "Right-click and choose Update Field"):
    """Insert a Word field code (TOC, PAGE, ...)."""
    r = paragraph.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin")
    r._r.append(fc)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def page_numbering(section, fmt: str, start: int | None = None):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))
    sectPr.append(el)


def footer_page_number(section, link_previous=False):
    section.footer.is_linked_to_previous = link_previous
    if link_previous:
        return
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    field(p, "PAGE", "1")
    for run in p.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(11)


# --- inline markdown --------------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")


def add_inline(paragraph, text: str, size=12, base_italic=False):
    text = text.replace("—", "—")
    for part in TOKEN.split(text):
        if not part:
            continue
        bold = italic = mono = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold = bold
        run.italic = italic or base_italic
        run.font.name = "Consolas" if mono else BODY_FONT
        run.font.size = Pt(size - 1 if mono else size)
    if not paragraph.runs:
        paragraph.add_run("")


def body_para(doc, text, style=None, size=12, justify=True, space_after=10):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(space_after)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, text, size=size)
    return p


def caption(doc, text, style_name):
    p = doc.add_paragraph(style=style_name)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    add_inline(p, text, size=10)
    return p


def add_table(doc, rows, tbl_caption=None):
    if tbl_caption:
        caption(doc, tbl_caption, "TableCaption")
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncols):
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            txt = row[j] if j < len(row) else ""
            add_inline(p, txt, size=10)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def bullet(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(4)
    add_inline(p, text, size=12)
    return p


def code_block(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.left_indent = Inches(0.3)
    pf.space_after = Pt(12)
    for i, ln in enumerate(lines):
        run = p.add_run(ln)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if i < len(lines) - 1:
            run.add_break()
    return p


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(16 if level <= 2 else 12)
    h.paragraph_format.space_after = Pt(8)
    h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = h.add_run(text)
    run.font.name = BODY_FONT
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = True
    run.font.size = Pt({1: 17, 2: 14, 3: 12.5}.get(level, 12))
    return h


# --- markdown block parser --------------------------------------------------
SKIP_LINE = re.compile(r"^\s*(\*Draft word count|\*\*Figures to insert:|<div|</div>)")


def render_markdown(doc, md: str, figures_by_section=None, drop_h1=True):
    md = re.sub(r"^>.*$", "", md, flags=re.M)          # draft note blocks
    lines = md.split("\n")
    i = 0
    current_section = None
    pending_table_caption = None

    def flush_figures():
        nonlocal current_section
        if figures_by_section and current_section in figures_by_section:
            for fn, cap in figures_by_section[current_section]:
                path = FIGDIR / fn
                if not path.exists():
                    path = SHOTDIR / fn          # screenshots of the artefact
                if not path.exists():
                    print(f"  MISSING FIGURE: {fn}")
                    continue
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.add_run().add_picture(str(path), width=Inches(5.6))
                caption(doc, cap, "FigureCaption")
            figures_by_section.pop(current_section, None)

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line or line == "---" or SKIP_LINE.match(line):
            i += 1
            continue

        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            code_block(doc, block)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1 and drop_h1:
                i += 1
                continue
            flush_figures()
            heading(doc, text, min(level, 3))
            if level == 2:
                current_section = text
            i += 1
            continue

        # a bold standalone line immediately before a table is its caption
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows, pending_table_caption)
            pending_table_caption = None
            continue

        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            bullet(doc, m.group(1))
            i += 1
            continue

        m = re.match(r"^\d+\.\s+(.*)$", line)
        if m:
            text = m.group(1)
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                text += " " + lines[i].strip()
                i += 1
            bullet(doc, text, numbered=True)
            continue

        # paragraph: gather until blank
        buf = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^\s*(\||#{1,4}\s|[-*]\s|\d+\.\s|```|---$)", lines[i]):
            if SKIP_LINE.match(lines[i].strip()):
                break
            buf.append(lines[i].strip())
            i += 1
        text = " ".join(buf)

        # Look past blank lines for the table. Markdown puts a blank line
        # between a caption and its table, and requiring the pipe on the very
        # next line meant no caption was ever recognised: every one was written
        # as body text, and the List of Tables field had nothing to collect.
        j = i
        while j < len(lines) and not lines[j].strip():
            j += 1
        nxt = lines[j].strip() if j < len(lines) else ""
        if re.fullmatch(r"\*\*Table [^*]+\*\*", text) and nxt.startswith("|"):
            pending_table_caption = text.strip("*")
            i = j
            continue
        body_para(doc, text)

    flush_figures()


# --- document ---------------------------------------------------------------
# The Arden cover template sits beside the repository, in the project folder.
COVER = DOCS.resolve().parent.parent / "Arden_Assessment_Cover_page (1).docx"

# What goes in the seven cells of the Arden cover page, in order.
COVER_FIELDS = [
    "Arden University",
    "MSc Cyber Security",
    "COM7014 Advanced Computing Project",
    "AI-Based Deepfake Detection Using Machine Learning with Robustness "
    "Evaluation",
    "STU1215887",
    "Dr. Muhammad Saqib",
    "{wordcount} (Chapters 1 to 8, excluding tables, references and appendices)",
]


def fill_cover(doc, wordcount: str) -> None:
    """Write the seven answers into the cover page's table.

    Each cell holds the printed label as an image in its FIRST paragraph and
    an empty paragraph beneath it for the answer. The answer therefore goes in
    paragraph 1: writing into paragraph 0 replaces the run that holds the
    label image, which silently deletes the label.
    """
    if not doc.tables:
        print("  WARNING: cover page has no table to fill")
        return
    for row, value in zip(doc.tables[0].rows, COVER_FIELDS):
        cell = row.cells[0]
        if len(cell.paragraphs) < 2:
            cell.add_paragraph()
        p = cell.paragraphs[1]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(2)
        pf.space_after = Pt(6)
        pf.left_indent = Inches(0.1)
        # The labels are floating images, so without a clearing break a short
        # answer rides up alongside its label instead of sitting under it.
        p.add_run().add_break(WD_BREAK.LINE_CLEAR_ALL)
        run = p.add_run(value.format(wordcount=wordcount))
        run.font.name = BODY_FONT
        run.font.size = Pt(12)
        run.bold = True

        # The seven rows overflow onto a second page at the document's default
        # spacing, so every paragraph in the cell is tightened: the label
        # images keep their own height, and the trailing empty paragraphs are
        # collapsed to nothing.
        for other in cell.paragraphs:
            if other is p:
                continue
            pf2 = other.paragraph_format
            pf2.space_before = Pt(0)
            pf2.space_after = Pt(0)
            pf2.line_spacing_rule = WD_LINE_SPACING.SINGLE
            if not other.runs and not other._p.findall(".//" + qn("w:drawing")):
                other.add_run().font.size = Pt(2)


def brand_header(section, logo_path) -> None:
    """Put a small Arden mark in a section's own header.

    The cover template's logo is about 30mm tall and sits behind the text, so
    carrying that header into the report overlaps the first two lines of every
    page. A smaller inline copy keeps the branding on every page without
    running into the body.
    """
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(logo_path), height=Mm(13))


def extract_cover_logo():
    """Pull the Arden mark out of the cover template's header, once."""
    import zipfile, tempfile
    if not COVER.exists():
        return None
    with zipfile.ZipFile(COVER) as z:
        names = [n for n in z.namelist() if n.startswith("word/media/")]
        if not names:
            return None
        # The header image is the only photograph among the label graphics.
        pick = max(names, key=lambda n: z.getinfo(n).file_size
                   if n.lower().endswith((".jpeg", ".jpg")) else -1)
        out = Path(tempfile.gettempdir()) / ("arden_mark" + Path(pick).suffix)
        out.write_bytes(z.read(pick))
        return out


def import_default_styles(doc) -> int:
    """Copy any style the blank python-docx template has and this one lacks.

    The Arden cover page was saved with a cut-down style set: it has no
    Heading, Caption or list styles at all, because the page never uses them.
    Building the report on top of it therefore fails until the usual
    definitions are put back.
    """
    from copy import deepcopy
    blank = Document()
    dst = doc.styles.element
    have = {s.get(qn("w:styleId")) for s in dst.findall(qn("w:style"))}
    added = 0
    for s in blank.styles.element.findall(qn("w:style")):
        sid = s.get(qn("w:styleId"))
        if sid in have:
            continue
        dst.append(deepcopy(s))
        added += 1
    return added


def front_heading(doc, text):
    """A chapter-sized heading that Word's TOC will not collect.

    Used only for the Table of Contents heading itself, which should not be
    an entry in its own contents list.
    """
    name = "FrontHeading"
    if name not in [s.name for s in doc.styles]:
        s = doc.styles.add_style(name, 1)
        s.base_style = doc.styles["Heading 1"]
        pPr = s.element.get_or_add_pPr()
        lvl = OxmlElement("w:outlineLvl")
        lvl.set(qn("w:val"), "9")          # body text: excluded from \o
        pPr.append(lvl)
    p = doc.add_paragraph(style=name)
    r = p.add_run(text)
    r.font.name = BODY_FONT
    return p


def main():
    doc = Document(str(COVER)) if COVER.exists() else Document()
    if not COVER.exists():
        print(f"  WARNING: cover page not found at {COVER}")
    else:
        n = import_default_styles(doc)
        print(f"cover page loaded; {n} missing styles imported")

    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    st.paragraph_format.space_after = Pt(10)

    # The Arden cover template carries a cut-down style set, so a base style
    # that exists in a blank python-docx document may be missing here.
    have = {s.name for s in doc.styles}
    for name in ("FigureCaption", "TableCaption"):
        s = doc.styles.add_style(name, 1)
        if "Caption" in have:
            s.base_style = doc.styles["Caption"]
        s.font.name = BODY_FONT
        s.font.size = Pt(10)
        s.font.italic = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        s.paragraph_format.space_after = Pt(12)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # A4, matching the Arden cover template exactly. The branding image in the
    # header is positioned against these margins, so changing them for the body
    # would shift the logo between the cover and the rest of the report.
    def page_setup(s):
        s.page_width = Mm(210)
        s.page_height = Mm(297)
        # Room for the branding mark in the header, which sits 16mm from the
        # page edge and is 13mm tall.
        s.top_margin = Mm(33)
        s.bottom_margin = Mm(25.4)
        s.left_margin = s.right_margin = Mm(25.4)
        s.header_distance = Mm(16)
        s.footer_distance = Mm(11)

    # ---------- cover page ----------
    # Section 0 is the Arden cover template, left exactly as supplied: its own
    # margins, its header image, and no page number. The table is filled in at
    # the end of this function, once the word count is known.

    # ---------- front matter: new section, roman from i ----------
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec)
    sec.different_first_page_header_footer = False
    # Own header carrying a smaller Arden mark, and own footer so the page
    # number does not also appear on the cover.
    logo = extract_cover_logo()
    if logo:
        brand_header(sec, logo)
    sec.footer.is_linked_to_previous = False
    page_numbering(sec, "lowerRoman", 1)
    footer_page_number(sec)

    fm = DOCS.joinpath("ch0_frontmatter.md").read_text(encoding="utf-8")

    def slice_between(text, start, end=None):
        a = text.index(start)
        b = text.index(end) if end else len(text)
        return text[a:b]

    # ---------- abstract ----------
    heading(doc, "Abstract", 1)
    abstract = slice_between(fm, "## Abstract", "## Acknowledgements")
    abstract = re.sub(r"^##.*$", "", abstract, flags=re.M)
    abstract = re.sub(r"^\*Target.*$", "", abstract, flags=re.M)
    render_markdown(doc, abstract)

    # ---------- acknowledgements ----------
    doc.add_page_break()
    heading(doc, "Acknowledgements", 1)
    ack = slice_between(fm, "## Acknowledgements", "## Abbreviations")
    ack = re.sub(r"^##.*$", "", ack, flags=re.M)
    render_markdown(doc, ack)

    # ---------- contents ----------
    doc.add_page_break()
    front_heading(doc, "Table of Contents")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    field(p, r'TOC \o "1-3" \h \z \u')

    doc.add_page_break()
    heading(doc, "List of Figures", 1)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    field(p, r'TOC \h \z \t "FigureCaption,1"')

    doc.add_page_break()
    heading(doc, "List of Tables", 1)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    field(p, r'TOC \h \z \t "TableCaption,1"')

    # ---------- abbreviations ----------
    doc.add_page_break()
    heading(doc, "Abbreviations", 1)
    abbr = slice_between(fm, "## Abbreviations")
    abbr = re.sub(r"^##.*$", "", abbr, flags=re.M)
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")]
            for ln in abbr.split("\n")
            if ln.strip().startswith("|") and not re.fullmatch(
                r"\|[\s:|-]+\|", ln.strip())]
    rows = [r for r in rows if any(r)]
    add_table(doc, [["Abbreviation", "Expansion"]] + rows)

    # ---------- body: new section, arabic from 1 ----------
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(body)
    page_numbering(body, "decimal", 1)
    footer_page_number(body)

    for idx, (fname, title) in enumerate(FILES):
        if idx:
            doc.add_page_break()
        heading(doc, title, 1)
        md = DOCS.joinpath(fname).read_text(encoding="utf-8")
        figs = {"ch6_results.md": FIGURES,
                "ch5_implementation.md": CH5_FIGURES,
                "ch4_design.md": CH4_FIGURES}.get(fname)
        render_markdown(doc, md, figures_by_section=figs)

    # ---------- references ----------
    doc.add_page_break()
    heading(doc, "References", 1)
    refs = DOCS.joinpath("references.md").read_text(encoding="utf-8")
    refs = re.sub(r"^#\s+References\s*$", "", refs, flags=re.M)
    for block in [b.strip() for b in refs.split("\n\n") if b.strip()]:
        text = " ".join(block.split())
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_after = Pt(10)
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)   # hanging indent
        add_inline(p, text, size=12)

    # ---------- appendices ----------
    doc.add_page_break()
    apx = DOCS.joinpath("appendices.md").read_text(encoding="utf-8")
    apx = re.sub(r"^#\s+Appendices\s*$", "", apx, flags=re.M)
    apx = apx.replace("\n# Appendix", "\n<<PB>>\n# Appendix")
    for n, chunk in enumerate(apx.split("<<PB>>")):
        if not chunk.strip():
            continue
        if n:
            doc.add_page_break()
        render_markdown(doc, chunk, drop_h1=False)

    # ---------- cover page word count ----------
    # Counted the way the limit is meant: Chapters 1 to 8 only, headings in,
    # tables, captions, references and appendices out.
    counting, done, words = False, False, 0
    for p in doc.paragraphs:
        text, style = p.text.strip(), (p.style.name or "")
        if not counting:
            if style.startswith("Heading 1") and text.lower().startswith("chapter 1"):
                counting = True
            else:
                continue
        if style.startswith("Heading 1") and (
                text.lower().startswith("reference") or
                text.lower().startswith("appendi")):
            done = True
        if done or "aption" in style:
            continue
        words += len(text.split())
    fill_cover(doc, f"{words:,} words")

    doc.save(OUT)
    print(f"saved {OUT}  ({OUT.stat().st_size/1e6:.2f} MB)")
    print(f"figures placed: {TOTAL_FIGURES}")
    print(f"main body word count on cover: {words:,}")


if __name__ == "__main__":
    main()
